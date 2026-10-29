# SPDX-License-Identifier: Apache-2.0
# SPDX-FileCopyrightText: 2026 William José Moreno Reyes

"""Regression tests for tenant-scoped evidence and optional report authorship."""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from flask import render_template_string
from flask_login import login_user
from openpyxl import load_workbook

from cacao_accounting import create_app
from cacao_accounting.database import (
    AccountingPeriod,
    ApprovalAction,
    ApprovalRequest,
    AuditTrail,
    Entity,
    Modules,
    Roles,
    RolesAccess,
    RolesUser,
    SalesOrder,
    User,
    UserCompanyAccess,
    database,
)
from cacao_accounting.document_authorship import AUTHOR_COLUMNS, batch_authorship, document_authorship, report_authorship
from cacao_accounting.reportes.internal_control import InternalControlFilters, get_internal_control_report
from cacao_accounting.reportes.services import PaginatedReport, ReportRow


@pytest.fixture()
def app():
    """Build two companies and independent auditor identities without external services."""
    instance = create_app(
        {"TESTING": True, "SQLALCHEMY_DATABASE_URI": "sqlite:///:memory:", "MODO_ESCRITORIO": False, "WTF_CSRF_ENABLED": False}
    )
    with instance.app_context():
        database.create_all()
        database.session.add_all(
            [
                Entity(code="A", name="A", company_name="A", tax_id="A", currency="NIO", enabled=True),
                Entity(code="B", name="B", company_name="B", tax_id="B", currency="NIO", enabled=True),
                Modules(id="accounting", module="accounting", default=True, enabled=True),
                User(id="admin", user="admin", name="Admin", active=True, password=b"x", classification="admin"),
                User(id="auditor", user="auditor", name="Auditor", active=True, password=b"x", classification="system"),
                Roles(id="audit-role", name="audit-role", note="Auditor"),
                RolesUser(user_id="auditor", role_id="audit-role", active=True),
                RolesAccess(rol_id="audit-role", module_id="accounting", access=True, report=True, view=True),
                UserCompanyAccess(user_id="auditor", company_code="A"),
                AccountingPeriod(
                    id="period-a",
                    entity="A",
                    name="2026-10",
                    enabled=True,
                    is_closed=False,
                    start=date(2026, 10, 1),
                    end=date(2026, 10, 31),
                ),
                SalesOrder(
                    id="order",
                    company="A",
                    document_no="SO-1",
                    docstatus=0,
                    created_by="auditor",
                    created=datetime(2026, 10, 1, 9),
                ),
            ]
        )
        database.session.commit()
        database.session.add_all(
            [
                AuditTrail(
                    id="event-a",
                    company="A",
                    document_type="sales_order",
                    document_id="order",
                    action="submitted",
                    actor_user_id="auditor",
                    actor_name="Auditor",
                    timestamp=datetime(2026, 10, 2, 10),
                    before_json='{"status":"draft"}',
                    after_json='{"status":"submitted"}',
                ),
                AuditTrail(
                    id="event-b",
                    company="B",
                    document_type="sales_order",
                    document_id="other",
                    action="submitted",
                    actor_user_id="admin",
                    timestamp=datetime(2026, 10, 2, 11),
                ),
                AuditTrail(
                    id="event-global",
                    document_type="roles",
                    document_id="audit-role",
                    action="submitted",
                    actor_user_id="admin",
                    timestamp=datetime(2026, 10, 2, 12),
                ),
                ApprovalRequest(
                    id="request",
                    document_type="sales_order",
                    document_id="order",
                    company_id="A",
                    requested_by="auditor",
                    status="Approved",
                    created_at=datetime(2026, 10, 1, 12),
                ),
                ApprovalAction(
                    id="action",
                    approval_request_id="request",
                    approved_by="admin",
                    action="approve",
                    created_at=datetime(2026, 10, 2, 10),
                ),
            ]
        )
        database.session.commit()
        yield instance
        database.session.remove()


def test_scoped_union_filters_and_pagination(app):
    """Audit and approvals share filters without exposing another tenant or global changes."""
    filters = InternalControlFilters(company="A", date_from=date(2026, 10, 2), date_to=date(2026, 10, 2))
    report = get_internal_control_report(filters, user_id="auditor")
    assert report.total_rows == 2
    assert {row.values["company"] for row in report.rows} == {"A"}
    assert {row.values["event_id"] for row in report.rows} == {"audit:event-a", "approval:action"}
    assert all(row.values["created_by"] == "Auditor" for row in report.rows)
    assert all(row.values["approved_by"] == "Admin" for row in report.rows)
    one = get_internal_control_report(
        InternalControlFilters(
            company="A", action="approved", actor_id="admin", document_type="sales_order", document_id="order", page_size=1
        ),
        user_id="admin",
    )
    assert one.total_rows == 1
    assert len(one.rows) == 1
    empty = get_internal_control_report(
        InternalControlFilters(company="A", action="submitted", page=2, page_size=1), user_id="auditor"
    )
    assert empty.rows == []
    assert empty.total_rows == 1


def test_global_access_and_period_bounds(app):
    """Global changes require Admin; period ownership and date consistency are enforced."""
    with pytest.raises(PermissionError):
        get_internal_control_report(InternalControlFilters(company="B"), user_id="auditor")
    with pytest.raises(PermissionError):
        get_internal_control_report(InternalControlFilters(company="A", include_global=True), user_id="auditor")
    report = get_internal_control_report(
        InternalControlFilters(company="A", include_global=True, action="submitted", period="period-a"), user_id="admin"
    )
    assert report.total_rows == 2
    with pytest.raises(ValueError):
        get_internal_control_report(InternalControlFilters(company="B", period="period-a"), user_id="admin")
    with pytest.raises(ValueError):
        get_internal_control_report(
            InternalControlFilters(company="A", period="period-a", date_from=date(2026, 10, 2)), user_id="admin"
        )
    with pytest.raises(ValueError):
        get_internal_control_report(
            InternalControlFilters(company="A", date_from=date(2026, 10, 3), date_to=date(2026, 10, 2)), user_id="admin"
        )
    for user_id in ("missing",):
        with pytest.raises(PermissionError):
            get_internal_control_report(InternalControlFilters(company="A"), user_id=user_id)


def test_authorship_flag_and_missing_evidence(app):
    """The boolean controls columns without replacing creator by ledger posting actor."""
    report = PaginatedReport(
        rows=[
            ReportRow(
                values={"voucher_type": "sales_order", "document_no": "SO-1", "created_by": "wrong-ledger-actor", "amount": 10}
            )
        ],
        totals={},
        columns=["voucher_type", "document_no", "created_by", "amount"],
    )
    hidden = report_authorship(report, "A", False)
    assert "created_by" not in hidden.columns
    shown = report_authorship(report, "A", True)
    assert set(AUTHOR_COLUMNS).issubset(shown.columns)
    assert shown.rows[0].values["created_by"] == "Auditor"
    assert shown.rows[0].values["created_at"] == datetime(2026, 10, 1, 9)
    assert shown.rows[0].values["approved_by"] == "Admin"
    assert shown.rows[0].values["approved_at"] == datetime(2026, 10, 2, 10)
    assert shown.rows[0].values["modified_by"] is None
    assert batch_authorship("B", {("sales_order", "order")}) == {}
    assert batch_authorship("A", {("unknown", "x")}) == {}
    unknown = report_authorship(
        PaginatedReport(rows=[ReportRow(values={"document_type": "unknown", "document_id": "x"})], totals={}), "A", True
    )
    assert unknown.rows[0].values["approved_by"] is None
    aggregate = PaginatedReport(rows=[ReportRow(values={"amount": 2})], totals={})
    assert report_authorship(aggregate, "A", True) is aggregate
    app.config["MODO_ESCRITORIO"] = True
    assert report_authorship(report, "A", True) is report
    assert batch_authorship("A", {("sales_order", "order")}) == {}
    with pytest.raises(PermissionError):
        get_internal_control_report(InternalControlFilters(company="A"), user_id="admin")


def test_detail_and_financial_edit_evidence(app):
    """Creation, approval and optional last modification are visible in document detail."""
    database.session.add(
        AuditTrail(
            company="A",
            document_type="sales_order",
            document_id="order",
            action="updated",
            actor_user_id="admin",
            actor_name="Admin",
            timestamp=datetime(2026, 10, 1, 11),
        )
    )
    database.session.commit()
    with app.test_request_context():
        login_user(database.session.get(User, "admin"))
        document = database.session.get(SalesOrder, "order")
        evidence = document_authorship("sales_order", document)
        assert evidence["modified_by"] == "Admin"
        html = render_template_string(
            '{% import "macros.html" as m %}{{ m.document_author_panel("sales_order", document) }}', document=document
        )
        for text in ("Creado por", "Aprobado por", "Modificado por", "Auditor", "Admin", "2026-10-02 10:00:00"):
            assert text in html


def test_route_and_exports_enforce_scope_and_preserve_evidence(app):
    """Exports match filtered queries and never expose another company's evidence."""
    client = app.test_client()
    with client.session_transaction() as session:
        session["_user_id"] = "auditor"
        session["_fresh"] = True
    url = "/reports/internal-control?company=A&action=submitted"
    response = client.get(url)
    assert response.status_code == 200
    assert b"audit:event-a" in response.data
    assert b"audit:event-b" not in response.data
    assert client.get(url + "&export=csv").data.count(b"audit:event-a") == 1
    excel = client.get(url + "&export=xlsx")
    assert excel.status_code == 200
    workbook = load_workbook(BytesIO(excel.data))
    assert any("audit:event-a" in row for row in workbook.active.values)
    assert client.get(url + "&include_global=1").status_code == 403
    assert client.get(url + "&date_from=bad").status_code == 400
    assert client.get(url + "&export=pdf").status_code == 400
    assert client.get("/reports/internal-control?company=B").status_code == 403
    app.config["MODO_ESCRITORIO"] = True
    assert client.get(url).status_code == 404


def test_review_provenance_and_legacy_missing_authors(app):
    """A review proposal has real preparation evidence even before its result exists."""
    database.session.add(
        ApprovalRequest(
            id="close-review",
            document_type="monthly_close",
            document_id="close",
            company_id="A",
            requested_by="auditor",
            status="Approved",
            created_at=datetime(2026, 10, 3, 8),
        )
    )
    database.session.flush()
    database.session.add_all(
        [
            ApprovalAction(
                approval_request_id="close-review", approved_by="admin", action="approve", created_at=datetime(2026, 10, 3, 10)
            ),
            AuditTrail(
                company="A",
                document_type="monthly_close",
                document_id="close",
                action="approved",
                actor_user_id="admin",
                actor_name="Admin",
                timestamp=datetime(2026, 10, 3, 10),
            ),
        ]
    )
    database.session.commit()
    evidence = batch_authorship("A", {("monthly_close", "close")})[("monthly_close", "close")]
    assert evidence["created_by"] == "Auditor"
    assert evidence["approved_by"] == "Admin"
    assert evidence["created_at"] == datetime(2026, 10, 3, 8)
    app.config["MODO_ESCRITORIO"] = True
    legacy = SalesOrder(id="legacy", company="A", docstatus=0)
    database.session.add(legacy)
    database.session.commit()
    app.config["MODO_ESCRITORIO"] = False
    assert document_authorship("sales_order", legacy)["created_by"] is None
    assert document_authorship("sales_order", legacy)["approved_by"] is None


def test_show_users_boolean_controls_view_and_export(app, monkeypatch):
    """The same boolean controls individual transaction columns in HTML and CSV."""
    from importlib import import_module

    routes = import_module("cacao_accounting.reportes.routes")

    report = PaginatedReport(
        rows=[
            ReportRow(
                values={
                    "voucher_type": "sales_order",
                    "document_no": "SO-1",
                    "debit": 10,
                    "created_by": "posting-actor",
                    "created_at": datetime(2026, 10, 2),
                }
            )
        ],
        totals={},
        columns=["voucher_type", "document_no", "debit", "created_by", "created_at"],
    )
    monkeypatch.setattr(routes, "get_account_movement_detail", lambda filters: report)
    client = app.test_client()
    with client.session_transaction() as session:
        session["_user_id"] = "admin"
        session["_fresh"] = True
    url = "/reports/account-movement?company=A&apply_filters=1"
    hidden = client.get(url + "&export=csv")
    assert b"approved_by" not in hidden.data
    assert b"created_by" not in hidden.data
    shown = client.get(url + "&show_users=1&export=csv")
    assert b"approved_by" in shown.data
    assert b"Auditor" in shown.data
    assert b"posting-actor" not in shown.data
    html = client.get(url + "&show_users=1")
    assert html.status_code == 200
    assert b'name="show_users"' in html.data
    assert b"Auditor" in html.data
    assert b"2026-10-02 10:00:00" in html.data


def test_spreadsheet_evidence_is_literal_and_inactive_user_is_denied(app):
    """User names cannot become executable spreadsheet formulas in evidence exports."""
    from cacao_accounting.reportes.helpers import _report_to_matrix

    report = PaginatedReport(rows=[ReportRow(values={"user": "=1+1", "amount": -10})], totals={}, columns=["user", "amount"])
    assert _report_to_matrix(report)[1] == [["'=1+1", -10]]
    user = database.session.get(User, "auditor")
    user.active = False
    database.session.commit()
    with pytest.raises(PermissionError):
        get_internal_control_report(InternalControlFilters(company="A"), user_id="auditor")


def test_xlsx_filter_values_are_written_as_literal_text(app):
    """Refs: 987; user-provided filters are not executable spreadsheet formulas."""
    from cacao_accounting.reportes.helpers import _write_financial_report_xlsx, _write_operational_report_xlsx
    from cacao_accounting.reportes.services import FinancialReportFilters

    formula_user = User(id="formula-user", user="=1+1", name="Formula User", active=True, password=b"x")
    database.session.add(formula_user)
    database.session.commit()
    report = PaginatedReport(rows=[], totals={"ending_balance": {"=1+1+": Decimal("100.00")}}, columns=[])
    with app.test_request_context("/"):
        login_user(formula_user)
        operational = _write_operational_report_xlsx(report, "operational", "Operational", {"company": "=1+1"})
        financial = _write_financial_report_xlsx(report, "financial", "Financial", FinancialReportFilters(company="=1+1"))

    for content in (operational, financial):
        workbook = load_workbook(BytesIO(content), data_only=False)
        cell = workbook["Filtros"]["B2"]
        assert cell.data_type == "s"
        assert cell.value == "'=1+1"
        user_cell = workbook.active["B3"]
        assert user_cell.data_type == "s"
        assert user_cell.value == "'=1+1"
        total_cell = workbook.active["C6"]
        assert total_cell.data_type == "s"
        assert total_cell.value == "'=1+1+ 100.00"
